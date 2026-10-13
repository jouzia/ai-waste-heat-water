"""Export Ali et al. (2020) supporting workbooks as source-linked raw row ledgers.

This is provenance-preserving extraction, not experimental interpretation or model
validation. It does not infer time units, flow direction, steady state, or operating
conditions from worksheet names. Original cell values and row locators are retained.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

SOURCE_DOI = "10.1371/journal.pone.0230207"
SOURCE_URL = "https://journals.plos.org/plosone/article?id=10.1371/journal.pone.0230207"
LICENSE = "CC BY 4.0"
WORKBOOKS = tuple(f"S{i}_Data.xlsx" for i in range(1, 6))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def serializable(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and not math.isfinite(value):
        return repr(value)
    return str(value)


def export_workbooks(input_dir: Path, output_dir: Path) -> dict[str, object]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install openpyxl>=3.1.2 to read source XLSX files") from exc

    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / "ali_2020_raw_observation_ledger.csv"
    manifest_path = output_dir / "ali_2020_raw_observation_ledger_manifest.json"
    records: list[dict[str, object]] = []
    workbook_manifest: list[dict[str, object]] = []
    sheet_manifest: list[dict[str, object]] = []

    for filename in WORKBOOKS:
        path = input_dir / filename
        if not path.is_file():
            raise FileNotFoundError(f"Required source workbook missing: {path}")
        source_hash = sha256_file(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        source_rows = 0
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            first = next(iterator, ())
            headers = [serializable(v).strip() for v in first]
            rows_here = 0
            for excel_row, row in enumerate(iterator, start=2):
                if not any(value is not None for value in row):
                    continue
                cells = {
                    f"column_{index + 1}": serializable(value)
                    for index, value in enumerate(row)
                }
                record = {
                    "source_doi": SOURCE_DOI,
                    "source_workbook": filename,
                    "source_sha256": source_hash,
                    "worksheet": sheet.title,
                    "excel_row_1based": excel_row,
                    "header_row_raw_json": json.dumps(headers, ensure_ascii=False),
                    "raw_cells_json": json.dumps(cells, ensure_ascii=False, sort_keys=True),
                }
                records.append(record)
                rows_here += 1
            source_rows += rows_here
            sheet_manifest.append(
                {
                    "workbook": filename,
                    "workbook_sha256": source_hash,
                    "worksheet": sheet.title,
                    "header_row_raw": headers,
                    "excel_max_row": sheet.max_row,
                    "excel_max_column": sheet.max_column,
                    "nonempty_data_rows_exported": rows_here,
                }
            )
        workbook.close()
        workbook_manifest.append(
            {
                "filename": filename,
                "size_bytes": path.stat().st_size,
                "sha256": source_hash,
                "nonempty_rows_exported": source_rows,
            }
        )

    fields = [
        "source_doi",
        "source_workbook",
        "source_sha256",
        "worksheet",
        "excel_row_1based",
        "header_row_raw_json",
        "raw_cells_json",
    ]
    with csv_path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)

    manifest: dict[str, object] = {
        "source_article_doi": SOURCE_DOI,
        "source_article_url": SOURCE_URL,
        "license": LICENSE,
        "license_url": "https://creativecommons.org/licenses/by/4.0/",
        "extracted_at_utc": datetime.now(UTC).isoformat(),
        "csv_filename": csv_path.name,
        "csv_sha256": sha256_file(csv_path),
        "total_nonempty_rows_exported": len(records),
        "sheets": sheet_manifest,
        "interpretation_policy": [
            "Values are preserved as source strings inside raw_cells_json; no unit conversions or imputations are performed.",
            "Worksheet names are source labels, not verified experimental condition metadata.",
            "Time units and the meaning/units of mw and FR are not inferred by this extractor.",
            "Repeated worksheets such as 80 and 80R remain separate source records.",
            "Blank rows are omitted; source worksheet and one-based Excel row locator are retained for every exported row.",
            "These data are not treated as held-out validation until source methods, condition mappings, model-input availability, and leakage risks are audited.",
        ],
        "workbooks": workbook_manifest,
    }
    manifest_path.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=Path("data/raw/ali_2020"))
    parser.add_argument("--output-dir", type=Path, default=Path("artifacts/ali_2020_exports"))
    args = parser.parse_args()
    manifest = export_workbooks(args.input_dir, args.output_dir)
    print(
        f"Exported {manifest['total_nonempty_rows_exported']} nonempty source rows "
        f"to {args.output_dir / manifest['csv_filename']}"
    )
    for workbook in manifest["workbooks"]:
        print(f"{workbook['filename']}: {workbook['nonempty_rows_exported']} rows; SHA-256 {workbook['sha256']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

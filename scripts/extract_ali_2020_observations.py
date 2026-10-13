"""Extract source-linked observations from Ali, Orfi & Najib (2020) supplements.

Outputs preserve raw workbook rows and distinguish timed raw mw readings from
temperature-only rows that lack timestamps. Mass-column semantics are audited
before any cumulative-mass or flux derivation is permitted.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from collections import defaultdict
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

SOURCE_DOI = "10.1371/journal.pone.0230207"
LICENSE = "CC BY 4.0"
MEMBRANE_AREA_M2 = 10.0
FLOW_BY_FILE = {
    "S2_Data.xlsx": 50,
    "S3_Data.xlsx": 100,
    "S4_Data.xlsx": 200,
    "S5_Data.xlsx": 300,
}
EXPECTED_SHA256 = {
    "S1_Data.xlsx": "e6e2d136a466c4c51ddc699bb751eb352df6154b9637f33581c8863a00c161ac",
    "S2_Data.xlsx": "cbca66c21d7eaf35e81ae77e753be8bd77a6db22f12025430f79fc8c069be010",
    "S3_Data.xlsx": "b6225af155380ed93f5ad51b31076adb0ea64c4fb33758222ef85c900969e3c1",
    "S4_Data.xlsx": "4897216f928ad63b53ed1a30e40c596a4c94e328c868ffb89291f7b3241bc428",
    "S5_Data.xlsx": "20ef92c9ea2a201af611aba74e1ed04c074acf3fce25db65f2ac8efdf028ab78",
}
EXPECTED_TRANSIENT_ROWS = 17_385
EXPECTED_COMPLETE_TIME_ROWS = 16_837
EXPECTED_TEMPERATURE_ONLY_ROWS = 548
EXPECTED_STEADY_ROWS = 16
STEADY_TARGETS_C = {50, 60, 70, 80}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]", "", str(value).strip().lower())


def finite_number(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    if not math.isfinite(float(value)):
        return None
    return float(value)


def _index(headers: list[str], *aliases: str) -> int | None:
    for alias in aliases:
        if alias in headers:
            return headers.index(alias)
    return None


def _cell(row: tuple[Any, ...], column: int | None) -> Any:
    return row[column] if column is not None and column < len(row) else None


def _raw_json(values: tuple[Any, ...]) -> str:
    return json.dumps(
        [None if value is None else str(value) for value in values],
        ensure_ascii=False,
    )


def _nonblank(values: tuple[Any, ...]) -> bool:
    return any(value is not None and str(value).strip() for value in values)


def _linear_slope(x: list[float], y: list[float]) -> float:
    if len(x) != len(y) or len(x) < 3:
        raise ValueError("At least three paired points are required for OLS.")
    x_bar = sum(x) / len(x)
    y_bar = sum(y) / len(y)
    denominator = sum((value - x_bar) ** 2 for value in x)
    if denominator <= 0:
        raise ValueError("Time values have no variation.")
    return sum(
        (x_value - x_bar) * (y_value - y_bar)
        for x_value, y_value in zip(x, y, strict=True)
    ) / denominator


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        raise ValueError(f"Refusing to write an empty observation file: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def _extract_steady_sheet(
    sheet: Any, *, filename: str, source_sha: str
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(
        (
            index
            for index, row in enumerate(rows)
            if "massflux" in [normalize_header(value) for value in row]
            and "volumeflowrate" in [normalize_header(value) for value in row]
        ),
        None,
    )
    if header_index is None:
        return [], [], {"sheet": sheet.title, "status": "steady_state_header_not_found"}

    headers = [normalize_header(value) for value in rows[header_index]]
    flow_column = _index(headers, "volumeflowrate")
    hot_in_column = _index(headers, "tein")
    hot_out_column = _index(headers, "teout")
    cold_in_column = _index(headers, "tcin", "tc")
    cold_out_column = _index(headers, "tcout")
    flux_column = _index(headers, "massflux")
    tds_hot_column = _index(headers, "tdse", "tds")
    tds_cold_column = _index(headers, "tdsc")
    sheet_flow_match = re.search(r"(\d+)", sheet.title)
    nominal_flow = int(sheet_flow_match.group(1)) if sheet_flow_match else None
    observations: list[dict[str, Any]] = []
    excluded: list[dict[str, Any]] = []

    for zero_index, row in enumerate(rows[header_index + 2 :], start=header_index + 2):
        if not _nonblank(row):
            continue
        flow = finite_number(_cell(row, flow_column))
        hot_in = finite_number(_cell(row, hot_in_column))
        flux = finite_number(_cell(row, flux_column))
        excel_row = zero_index + 1
        if (
            flow is None
            or hot_in is None
            or flux is None
            or nominal_flow is None
            or abs(flow - nominal_flow) > 1.0
            or round(hot_in) not in STEADY_TARGETS_C
        ):
            excluded.append(
                {
                    "supporting_data_id": "S1",
                    "workbook_filename": filename,
                    "worksheet_name": sheet.title,
                    "excel_row_1based": excel_row,
                    "reason": "not_a_valid_flow_temperature_mass_flux_condition",
                    "raw_row_json": _raw_json(row),
                }
            )
            continue
        observations.append(
            {
                "observation_id": f"ALI2020-S1-{sheet.title}-row{excel_row}",
                "source_article_doi": SOURCE_DOI,
                "supporting_data_id": "S1",
                "workbook_filename": filename,
                "source_file_sha256": source_sha,
                "worksheet_name": sheet.title,
                "excel_row_1based": excel_row,
                "nominal_flow_l_h": flow,
                "hot_inlet_temperature_c": hot_in,
                "hot_outlet_temperature_c": _cell(row, hot_out_column),
                "cold_inlet_temperature_c": _cell(row, cold_in_column),
                "cold_outlet_temperature_c": _cell(row, cold_out_column),
                "mass_flux_kg_m2_h": flux,
                "tds_hot_ppm_raw": _cell(row, tds_hot_column),
                "tds_cold_ppm_raw": _cell(row, tds_cold_column),
                "raw_row_json": _raw_json(row),
            }
        )
    return observations, excluded, {
        "sheet": sheet.title,
        "header_excel_row": header_index + 1,
        "steady_state_observation_rows": len(observations),
        "excluded_noncondition_rows": len(excluded),
    }


def _extract_transient_sheet(
    sheet: Any, *, filename: str, source_sha: str, nominal_flow: int
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(
        (
            index
            for index, row in enumerate(rows)
            if "time" in [normalize_header(value) for value in row]
            and {"mw", "mwkg"}.intersection(
                normalize_header(value) for value in row
            )
        ),
        None,
    )
    if header_index is None:
        return [], {"sheet": sheet.title, "status": "time_series_header_not_found"}

    header_row = rows[header_index]
    headers = [normalize_header(value) for value in header_row]
    time_column = _index(headers, "time")
    mass_column = _index(headers, "mw", "mwkg")
    flow_column = _index(headers, "fr", "frlh")
    hot_in_column = _index(headers, "thin", "thim", "tein")
    hot_out_column = _index(headers, "thout", "teout", "thut")
    cold_in_column = _index(headers, "tcin")
    cold_out_column = _index(headers, "tcout")

    match = re.fullmatch(r"(\d+)(R?)", sheet.title.strip(), flags=re.IGNORECASE)
    target_temperature = int(match.group(1)) if match else None
    replicate_label = "R" if match and match.group(2).upper() == "R" else ""
    observations: list[dict[str, Any]] = []
    complete_count = 0
    temperature_only_count = 0
    other_incomplete_count = 0

    for zero_index, row in enumerate(rows[header_index + 1 :], start=header_index + 1):
        if not _nonblank(row):
            continue
        time_value = finite_number(_cell(row, time_column))
        mass_value = finite_number(_cell(row, mass_column))
        hot_in = _cell(row, hot_in_column)
        if time_value is not None and mass_value is not None:
            row_status = "complete_time_series_row"
            complete_count += 1
        elif time_value is None and mass_value is None and finite_number(hot_in) is not None:
            row_status = "temperature_only_no_timestamp"
            temperature_only_count += 1
        else:
            row_status = "incomplete_or_metadata_row"
            other_incomplete_count += 1

        excel_row = zero_index + 1
        observations.append(
            {
                "observation_id": (
                    f"ALI2020-{Path(filename).stem}-{sheet.title}-row{excel_row}"
                ),
                "source_article_doi": SOURCE_DOI,
                "supporting_data_id": Path(filename).stem.replace("_Data", ""),
                "workbook_filename": filename,
                "source_file_sha256": source_sha,
                "worksheet_name": sheet.title,
                "excel_row_1based": excel_row,
                "nominal_flow_l_h": nominal_flow,
                "hot_inlet_target_c": target_temperature,
                "replicate_label": replicate_label,
                "time_min": time_value,
                "hot_inlet_temperature_c": hot_in,
                "hot_outlet_temperature_c": _cell(row, hot_out_column),
                "cold_inlet_temperature_c": _cell(row, cold_in_column),
                "cold_outlet_temperature_c": _cell(row, cold_out_column),
                "cumulative_distillate_mass_kg": None,
                "mw_raw_reading_kg": mass_value,
                "mw_semantics_status": "unverified_raw_mw_measurement",
                "fr_raw_reading": _cell(row, flow_column),
                "fr_units_label_from_header": "L/h" if flow_column is not None else None,
                "fr_semantics_status": "unverified_role_relative_to_nominal_flow",
                "row_status": row_status,
                "raw_header_json": _raw_json(header_row),
                "raw_row_json": _raw_json(row),
            }
        )
    return observations, {
        "sheet": sheet.title,
        "header_excel_row": header_index + 1,
        "nonblank_rows": len(observations),
        "complete_time_series_rows": complete_count,
        "temperature_only_no_timestamp_rows": temperature_only_count,
        "other_incomplete_rows": other_incomplete_count,
        "time_unit": "minute; source paper reports 10-second sampling",
        "nominal_flow_l_h": nominal_flow,
        "hot_inlet_target_c": target_temperature,
        "replicate_label": replicate_label,
    }


def _mw_semantics_audit(
    transient_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Summarize raw mw behavior without assuming it is cumulative mass."""
    grouped: dict[tuple[str, int, str], list[dict[str, Any]]] = defaultdict(list)
    for row in transient_rows:
        if row["row_status"] == "complete_time_series_row":
            grouped[
                (
                    str(row["workbook_filename"]),
                    int(row["nominal_flow_l_h"]),
                    str(row["worksheet_name"]),
                )
            ].append(row)

    output: list[dict[str, Any]] = []
    for (filename, flow, sheet_name), rows in sorted(grouped.items()):
        rows = sorted(rows, key=lambda row: float(row["time_min"]))
        time_values = [float(row["time_min"]) for row in rows]
        mw_values = [float(row["mw_raw_reading_kg"]) for row in rows]
        differences = [right - left for left, right in zip(mw_values, mw_values[1:])]
        negative = [difference for difference in differences if difference < 0.0]
        positive = [difference for difference in differences if difference > 0.0]
        target_match = re.fullmatch(r"(\d+)(R?)", sheet_name.strip(), flags=re.IGNORECASE)
        output.append(
            {
                "workbook_filename": filename,
                "source_file_sha256": rows[0]["source_file_sha256"],
                "nominal_flow_l_h_crosswalk": flow,
                "worksheet_name": sheet_name,
                "hot_inlet_target_c_from_sheet_label": (
                    int(target_match.group(1)) if target_match else None
                ),
                "replicate_label": rows[0]["replicate_label"],
                "row_count": len(rows),
                "time_start_min": float(time_values[0]),
                "time_end_min": float(time_values[-1]),
                "mw_raw_start_kg": float(mw_values[0]),
                "mw_raw_end_kg": float(mw_values[-1]),
                "mw_raw_min_kg": float(min(mw_values)),
                "mw_raw_max_kg": float(max(mw_values)),
                "first_positive_mw_time_min": (
                    float(time_values[next(index for index, value in enumerate(mw_values) if value > 0.0)])
                    if any(value > 0.0 for value in mw_values)
                    else None
                ),
                "negative_step_count": int(len(negative)),
                "positive_step_count": int(len(positive)),
                "max_negative_step_magnitude_kg": (
                    float(abs(min(negative))) if negative else 0.0
                ),
                "median_negative_step_magnitude_kg": (
                    float(sorted(abs(value) for value in negative)[len(negative) // 2]) if negative else 0.0
                ),
                "zero_fraction": sum(value == 0.0 for value in mw_values) / len(mw_values),
                "raw_mw_monotonic": all(difference >= 0.0 for difference in differences),
                "status": (
                    "nonmonotonic_raw_mw_semantics_unverified"
                    if any(difference < 0.0 for difference in differences)
                    else "monotonic_raw_mw_semantics_still_unverified"
                ),
                "interpretation": (
                    "Raw mw readings are preserved without cumulative-mass interpretation. "
                    "Do not derive cumulative production or flux from this column until "
                    "the source measurement/reset protocol is verified."
                ),
            }
        )
    return output


def extract_all(input_dir: Path, output_dir: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install openpyxl>=3.1.2 to read the supporting XLSX files.") from exc

    if not input_dir.is_dir():
        raise FileNotFoundError(f"Supporting-data directory not found: {input_dir}")
    output_dir.mkdir(parents=True, exist_ok=True)
    transient_rows: list[dict[str, Any]] = []
    steady_rows: list[dict[str, Any]] = []
    excluded_steady_rows: list[dict[str, Any]] = []
    file_records: list[dict[str, Any]] = []

    for filename in ("S1_Data.xlsx", "S2_Data.xlsx", "S3_Data.xlsx", "S4_Data.xlsx", "S5_Data.xlsx"):
        path = input_dir / filename
        if not path.is_file():
            raise FileNotFoundError(f"Required supporting workbook is missing: {path}")
        source_sha = sha256_file(path)
        if source_sha != EXPECTED_SHA256[filename]:
            raise ValueError(
                f"Source hash changed for {filename}; audit before extraction. "
                f"Expected {EXPECTED_SHA256[filename]}, got {source_sha}."
            )
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet_records: list[dict[str, Any]] = []
        if filename == "S1_Data.xlsx":
            for sheet in workbook.worksheets:
                extracted, excluded, sheet_record = _extract_steady_sheet(
                    sheet, filename=filename, source_sha=source_sha
                )
                steady_rows.extend(extracted)
                excluded_steady_rows.extend(excluded)
                sheet_records.append(sheet_record)
        else:
            nominal_flow = FLOW_BY_FILE[filename]
            for sheet in workbook.worksheets:
                extracted, sheet_record = _extract_transient_sheet(
                    sheet,
                    filename=filename,
                    source_sha=source_sha,
                    nominal_flow=nominal_flow,
                )
                transient_rows.extend(extracted)
                sheet_records.append(sheet_record)
        workbook.close()
        file_records.append(
            {
                "filename": filename,
                "size_bytes": path.stat().st_size,
                "sha256": source_sha,
                "supporting_data_id": filename.removesuffix("_Data.xlsx"),
                "sheets": sheet_records,
            }
        )

    complete_rows = [
        row for row in transient_rows if row["row_status"] == "complete_time_series_row"
    ]
    temperature_only_rows = [
        row
        for row in transient_rows
        if row["row_status"] == "temperature_only_no_timestamp"
    ]
    other_incomplete_rows = [
        row for row in transient_rows if row["row_status"] == "incomplete_or_metadata_row"
    ]
    if (
        len(transient_rows) != EXPECTED_TRANSIENT_ROWS
        or len(complete_rows) != EXPECTED_COMPLETE_TIME_ROWS
        or len(temperature_only_rows) != EXPECTED_TEMPERATURE_ONLY_ROWS
        or other_incomplete_rows
        or len(steady_rows) != EXPECTED_STEADY_ROWS
    ):
        raise ValueError(
            "Extracted row counts differ from the audited source version: "
            f"transient={len(transient_rows)}, complete={len(complete_rows)}, "
            f"temperature-only={len(temperature_only_rows)}, "
            f"other-incomplete={len(other_incomplete_rows)}, steady={len(steady_rows)}."
        )

    mw_audit_rows = _mw_semantics_audit(transient_rows)
    transient_path = output_dir / "ali_2020_transient_observation_ledger.csv"
    steady_path = output_dir / "ali_2020_steady_state_observation_ledger.csv"
    excluded_path = output_dir / "ali_2020_excluded_steady_state_rows.csv"
    mw_audit_path = output_dir / "ali_2020_mw_semantics_audit.csv"
    _write_csv(transient_path, transient_rows)
    _write_csv(steady_path, steady_rows)
    if excluded_steady_rows:
        _write_csv(excluded_path, excluded_steady_rows)
    _write_csv(mw_audit_path, mw_audit_rows)

    manifest = {
        "source_article_doi": SOURCE_DOI,
        "source_article_url": (
            "https://journals.plos.org/plosone/article?id=10.1371/journal.pone.0230207"
        ),
        "license": LICENSE,
        "license_url": "https://creativecommons.org/licenses/by/4.0/",
        "extracted_at_utc": datetime.now(UTC).isoformat(),
        "flow_crosswalk": FLOW_BY_FILE,
        "flow_crosswalk_basis": (
            "S2-S5 are mapped to nominal 50/100/200/300 L/h conditions from the "
            "paper's operating matrix and workbook organization. FR is retained as a "
            "raw column and is not used as a model input because its role relative to "
            "the nominal circulation flow is not verified."
        ),
        "time_unit": "minute",
        "sampling_interval_reported_by_source": "10 seconds",
        "mw_raw_reading_unit_label": "kg (from source column header)",
        "mass_semantics_status": "unverified; raw mw is nonmonotonic in multiple traces",
        "expected_source_sha256": EXPECTED_SHA256,
        "files": file_records,
        "outputs": {
            "transient_ledger": transient_path.name,
            "transient_rows": len(transient_rows),
            "complete_time_series_rows": len(complete_rows),
            "temperature_only_no_timestamp_rows": len(temperature_only_rows),
            "other_incomplete_rows": len(other_incomplete_rows),
            "steady_state_ledger": steady_path.name,
            "steady_state_rows": len(steady_rows),
            "excluded_steady_state_rows": excluded_path.name,
            "excluded_steady_state_row_count": len(excluded_steady_rows),
            "mw_semantics_audit": mw_audit_path.name,
            "mw_semantics_audit_rows": len(mw_audit_rows),
            "tail_flux_consistency_check": "disabled_pending_mw_semantics_verification",
            "tail_flux_check_rows": 0,
        },
        "data_quality_findings": [
            "S5 50/60/70 sheets contain temperature-only rows without timestamps after the timed series; retained as temperature_only_no_timestamp and excluded from time-domain fitting.",
            "S2 55 begins with a nonzero raw mw reading and its final measured hot-inlet temperature is below the sheet label; treat as a partial/nonstandard trace unless source documentation resolves the condition.",
            "S5 80 ends at a measured hot-inlet temperature below the 80 C sheet label; do not treat it as a fully attained 80 C steady-state condition.",
            "The S1 Q=300 workbook contains a numeric row that is not a valid operating condition; it is excluded and archived in the excluded-row ledger.",
            "Raw mw is nonmonotonic in multiple transient traces; cumulative-mass and tail-slope flux comparisons are disabled until source measurement/reset semantics are verified.",
        ],
        "transformation_policy": [
            "All raw time-series values and source row locators are preserved; no smoothing or interpolation is applied.",
            "Time is retained in source minutes; raw mw readings are retained in kg with no cumulative-mass interpretation until verified.",
            "FR is retained as a raw source field; its role relative to nominal circulation flow is not inferred.",
            "Temperature-only rows without timestamps remain in the ledger but are excluded from time-domain fitting.",
            "Source hashes are pinned; changed source bytes require manual provenance review before extraction.",
        ],
    }
    manifest_path = output_dir / "ali_2020_observation_ledger_manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(
        "Extracted "
        f"{len(transient_rows)} transient rows ({len(complete_rows)} timed), "
        f"{len(temperature_only_rows)} temperature-only rows, "
        f"{len(steady_rows)} steady-state observations."
    )
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=Path("data/raw/ali_2020"))
    parser.add_argument(
        "--output-dir", type=Path, default=Path("artifacts/ali_2020_exports")
    )
    args = parser.parse_args()
    extract_all(args.input_dir, args.output_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

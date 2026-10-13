"""Download and inventory Ali, Orfi & Najib (2020) CC BY 4.0 DCMD supporting XLSX files.

The script preserves source bytes, checks ZIP/XLSX signatures, records SHA-256,
and inventories workbook structure. It does not promote values to observations;
a separate source-linked extraction and validation step is required.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import time
from datetime import UTC, datetime
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

SOURCE_DOI = "10.1371/journal.pone.0230207"
SOURCE_PAGE = "https://journals.plos.org/plosone/article?id=10.1371/journal.pone.0230207"
LICENSE_URL = "https://creativecommons.org/licenses/by/4.0/"
DATA_OBJECTS = tuple(
    {
        "supporting_data_id": f"S{i}",
        "doi": f"10.1371/journal.pone.0230207.s{suffix}",
        "filename": f"S{i}_Data.xlsx",
        "url": (
            "https://journals.plos.org/plosone/article/file"
            f"?id=10.1371/journal.pone.0230207.s{suffix}&type=supplementary"
        ),
    }
    for i, suffix in enumerate(("002", "003", "004", "005", "006"), start=1)
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def download_xlsx(url: str, target: Path, *, force: bool = False) -> dict[str, object]:
    if target.exists() and not force:
        prefix = target.read_bytes()[:4]
        if prefix[:2] != b"PK":
            raise ValueError(f"Existing file is not an XLSX/ZIP object: {target}")
        return {
            "status": "already_present",
            "size_bytes": target.stat().st_size,
            "sha256": sha256_file(target),
        }

    partial = target.with_suffix(target.suffix + ".part")
    request = Request(
        url,
        headers={"User-Agent": "ai-waste-heat-water/1.0 (research data provenance)"},
    )
    last_error: Exception | None = None
    for attempt in range(4):
        try:
            with urlopen(request, timeout=120) as response, partial.open("wb") as output:
                while chunk := response.read(1024 * 1024):
                    output.write(chunk)
            if partial.read_bytes()[:2] != b"PK":
                raise ValueError(
                    f"Source endpoint did not return an XLSX/ZIP payload: {url}"
                )
            partial.replace(target)
            return {
                "status": "downloaded",
                "size_bytes": target.stat().st_size,
                "sha256": sha256_file(target),
            }
        except (HTTPError, URLError, TimeoutError, OSError, ValueError) as exc:
            last_error = exc
            partial.unlink(missing_ok=True)
            if attempt < 3:
                time.sleep(min(2**attempt, 20))
    raise RuntimeError(f"Could not download a valid XLSX from {url}: {last_error}") from last_error


def inspect_xlsx(path: Path) -> list[dict[str, object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install workbook inspection support with openpyxl>=3.1.2") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[dict[str, object]] = []
    for sheet in workbook.worksheets:
        preview = []
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 0, 5),
                                   min_col=1, max_col=min(sheet.max_column or 0, 12),
                                   values_only=True):
            preview.append([None if value is None else str(value) for value in row])
        sheets.append(
            {
                "sheet_name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "preview_first_5_rows_first_12_columns": preview,
            }
        )
    workbook.close()
    return sheets


def acquire_all(output_dir: Path, manifest_path: Path, *, force: bool = False) -> dict[str, object]:
    output_dir.mkdir(parents=True, exist_ok=True)
    records: list[dict[str, object]] = []
    for item in DATA_OBJECTS:
        target = output_dir / item["filename"]
        result = download_xlsx(item["url"], target, force=force)
        sheets = inspect_xlsx(target)
        records.append(
            {
                **item,
                "local_path": str(target),
                **result,
                "status": "downloaded_and_inspected",
                "sheets": sheets,
            }
        )

    manifest: dict[str, object] = {
        "source_article_doi": SOURCE_DOI,
        "source_article_url": SOURCE_PAGE,
        "license": "CC BY 4.0",
        "license_url": LICENSE_URL,
        "retrieved_at_utc": datetime.now(UTC).isoformat(),
        "files": records,
        "scope_warning": (
            "Workbook inventory and previews only. The files are source data, but "
            "no values have been promoted to the repository's observation ledger. "
            "Review each sheet's variables, units, experimental conditions, and "
            "measurement provenance before quantitative model validation."
        ),
    }
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=Path("data/raw/ali_2020"))
    parser.add_argument(
        "--manifest",
        type=Path,
        default=Path("artifacts/provenance/ali_2020_supporting_data_manifest.json"),
    )
    parser.add_argument("--force", action="store_true", help="Redownload existing workbooks.")
    args = parser.parse_args()
    manifest = acquire_all(args.output_dir, args.manifest, force=args.force)
    for item in manifest["files"]:
        print(
            f"{item['supporting_data_id']}: {item['status']} "
            f"{item['size_bytes']} bytes SHA-256={item['sha256']}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
